"""./lect/CH-0N-Content.md 강의안을 같은 이름의 .pptx 파일로 변환한다.

각 마크다운 페이지(--- 로 구분)를 슬라이드 한 장으로 만들고,
'> **[Note]**' 블록은 화면에 표시되지 않는 PowerPoint 발표자 노트로 옮긴다.
"""
import glob
import os
import re

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_CONNECTOR
from PIL import Image

FONT_NAME = "맑은 고딕"
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
MARGIN_IN = 0.5
CONTENT_W_IN = SLIDE_W_IN - 2 * MARGIN_IN

TITLE_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
TEXT_COLOR = RGBColor(0x22, 0x22, 0x22)
CODE_BG = RGBColor(0xF2, 0xF2, 0xF2)
CODE_TEXT = RGBColor(0x33, 0x33, 0x33)
TABLE_HEADER_BG = RGBColor(0x1F, 0x3A, 0x5F)

PAGE_HEADER_RE = re.compile(r"^#{1,3}\s*페이지\s*(\d+)\.\s*(.*)$")
NOTE_RE = re.compile(r"^>\s*\*\*\[Note\]\*\*\s*(.*)$")
BULLET_RE = re.compile(r"^(\s*)-\s+(.*)$")
NUM_RE = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
TABLE_ROW_RE = re.compile(r"^\|.*\|$")
TABLE_SEP_RE = re.compile(r"^\|[\s:|-]+\|$")
IMAGE_RE = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)$")


def strip_inline_code(text: str) -> str:
    return text.replace("`", "")


def parse_pages(md_text: str):
    lines = md_text.split("\n")
    blocks, current = [], []
    for line in lines:
        if line.strip() == "---":
            blocks.append(current)
            current = []
        else:
            current.append(line)
    blocks.append(current)

    pages = []
    for block in blocks:
        page = parse_block(block)
        if page:
            pages.append(page)
    return pages


def parse_block(lines):
    title_idx = None
    for i, line in enumerate(lines):
        if line.startswith("#"):
            title_idx = i
            break
    if title_idx is None:
        return None

    m = PAGE_HEADER_RE.match(lines[title_idx].strip())
    top_level = lines[title_idx].startswith("# ") and not lines[title_idx].startswith("### ")
    if m:
        page_num, title = m.group(1), m.group(2).strip()
    else:
        page_num, title = "", lines[title_idx].lstrip("#").strip()

    body_lines = lines[title_idx + 1:]
    note_text, body_lines = extract_note(body_lines)
    segments = parse_segments(body_lines)
    return {
        "page_num": page_num,
        "title": strip_inline_code(title),
        "top_level": top_level,
        "segments": segments,
        "note": note_text,
    }


def extract_note(lines):
    note = None
    kept = []
    for line in lines:
        m = NOTE_RE.match(line.strip())
        if m:
            note = m.group(1).strip()
        else:
            kept.append(line)
    return note, kept


def parse_segments(lines):
    segments = []
    pending = []
    i = 0
    n = len(lines)

    def flush_text():
        if pending:
            segments.append({"type": "text", "items": list(pending)})
            pending.clear()

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_text()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            segments.append({"type": "code", "text": "\n".join(code_lines)})
            continue

        if TABLE_ROW_RE.match(stripped):
            flush_text()
            table_rows = []
            while i < n and TABLE_ROW_RE.match(lines[i].strip()):
                row_line = lines[i].strip()
                if not TABLE_SEP_RE.match(row_line):
                    cells = [c.strip() for c in row_line.strip("|").split("|")]
                    table_rows.append(cells)
                i += 1
            if table_rows:
                segments.append({"type": "table", "rows": table_rows})
            continue

        if stripped == "":
            i += 1
            continue

        im = IMAGE_RE.match(stripped)
        if im:
            flush_text()
            segments.append({"type": "image", "path": im.group(1).strip()})
            i += 1
            continue

        bm = BULLET_RE.match(line)
        nm = NUM_RE.match(line)
        if bm:
            level = min(len(bm.group(1)) // 2, 1)
            pending.append({"kind": "bullet", "level": level, "text": strip_inline_code(bm.group(2))})
        elif nm:
            pending.append({"kind": "plain", "level": 0, "text": strip_inline_code(f"{nm.group(2)}. {nm.group(3)}")})
        else:
            pending.append({"kind": "plain", "level": 0, "text": strip_inline_code(stripped)})
        i += 1

    flush_text()
    return segments


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------

def _chars_per_line(width_in, font_pt):
    # Korean/CJK characters render roughly square (width ~= point size), so a
    # box `width_in` inches wide fits about width_in*72/font_pt full-width
    # characters per line. 1.05 gives a little slack for mixed Latin/CJK text.
    return max(10, int(width_in * 72 / (font_pt * 1.05)))


def estimate_text_height(segment, width_in, font_pt):
    total_lines = 0
    for item in segment["items"]:
        item_font_pt = font_pt if item["level"] == 0 else font_pt - 2
        chars_per_line = _chars_per_line(width_in - item["level"] * 0.3, item_font_pt)
        text = ("• " if item["kind"] == "bullet" else "") + item["text"]
        total_lines += max(1, -(-len(text) // chars_per_line))
    line_h_in = (font_pt + 6) / 72
    return max(0.4, total_lines * line_h_in)


def estimate_code_height(segment, font_pt=11):
    lines = segment["text"].split("\n")
    line_h_in = (font_pt + 3) / 72
    return max(0.35, len(lines) * line_h_in) + 0.15


def estimate_table_height(segment, font_pt=12):
    line_h_in = (font_pt + 14) / 72
    return max(0.5, len(segment["rows"]) * line_h_in)


def estimate_image_size(segment, base_dir, max_w_in, max_h_in=4.2):
    abs_path = os.path.join(base_dir, segment["path"])
    try:
        with Image.open(abs_path) as im:
            px_w, px_h = im.size
    except (OSError, FileNotFoundError):
        return max_w_in, 2.0
    aspect = px_h / px_w
    w_in = max_w_in
    h_in = w_in * aspect
    if h_in > max_h_in:
        h_in = max_h_in
        w_in = h_in / aspect
    return w_in, h_in


def add_text_segment(slide, segment, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    first = True
    for item in segment["items"]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = item["level"]
        prefix = "• " if item["kind"] == "bullet" else ""
        run = p.add_run()
        run.text = prefix + item["text"]
        run.font.size = Pt(16 if item["level"] == 0 else 14)
        run.font.name = FONT_NAME
        run.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(6)
    return box


def add_code_segment(slide, segment, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = CODE_BG
    box.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)
    box.line.width = Pt(0.75)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    lines = segment["text"].split("\n")
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = line if line.strip() else " "
        run.font.size = Pt(11)
        run.font.name = FONT_NAME
        run.font.color.rgb = CODE_TEXT
        p.space_after = Pt(0)
    return box


def add_table_segment(slide, segment, x, y, w, h):
    rows_data = segment["rows"]
    n_rows = len(rows_data)
    n_cols = max(len(r) for r in rows_data)
    rows_data = [r + [""] * (n_cols - len(r)) for r in rows_data]

    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    col_w = Emu(int(Inches(w) / n_cols))
    for c in range(n_cols):
        table.columns[c].width = col_w

    for r, row in enumerate(rows_data):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = strip_inline_code(cell_text)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.name = FONT_NAME
                if r == 0:
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    p.font.color.rgb = TEXT_COLOR
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_HEADER_BG
    return shape


def add_image_segment(slide, segment, base_dir, x, y, w_in, h_in):
    abs_path = os.path.join(base_dir, segment["path"])
    if not os.path.isfile(abs_path):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w_in), Inches(0.4))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = f"[이미지 없음: {segment['path']}]"
        run.font.size = Pt(12)
        run.font.name = FONT_NAME
        run.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
        return box
    return slide.shapes.add_picture(abs_path, Inches(x), Inches(y), width=Inches(w_in), height=Inches(h_in))


def render_section_slide(prs, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1.0), Inches(2.7), Inches(SLIDE_W_IN - 2.0), Inches(2.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = page["title"]
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.name = FONT_NAME
    run.font.color.rgb = TITLE_COLOR
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return slide


def render_content_slide(prs, page, base_dir):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(MARGIN_IN), Inches(0.35), Inches(CONTENT_W_IN), Inches(0.9))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = page["title"]
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.name = FONT_NAME
    run.font.color.rgb = TITLE_COLOR

    line_box = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(MARGIN_IN), Inches(1.25), Inches(SLIDE_W_IN - MARGIN_IN), Inches(1.25)
    )
    line_box.line.color.rgb = RGBColor(0xC8, 0xC8, 0xC8)
    line_box.line.width = Pt(1)

    cursor_y = 1.45
    max_y = SLIDE_H_IN - 0.2
    for segment in page["segments"]:
        if segment["type"] == "text":
            h = estimate_text_height(segment, CONTENT_W_IN, 16)
            add_text_segment(slide, segment, MARGIN_IN, cursor_y, CONTENT_W_IN, h)
        elif segment["type"] == "code":
            h = estimate_code_height(segment)
            add_code_segment(slide, segment, MARGIN_IN, cursor_y, CONTENT_W_IN, h)
        elif segment["type"] == "table":
            h = estimate_table_height(segment)
            add_table_segment(slide, segment, MARGIN_IN, cursor_y, CONTENT_W_IN, h)
        elif segment["type"] == "image":
            w, h = estimate_image_size(segment, base_dir, CONTENT_W_IN)
            x = MARGIN_IN + (CONTENT_W_IN - w) / 2
            add_image_segment(slide, segment, base_dir, x, cursor_y, w, h)
        else:
            h = 0
        cursor_y += h + 0.15

    return slide


def build_pptx(md_path: str, out_path: str):
    with open(md_path, encoding="utf-8") as f:
        text = f.read()
    pages = parse_pages(text)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)

    base_dir = os.path.dirname(os.path.abspath(md_path))
    for page in pages:
        if page["top_level"] and not page["segments"]:
            slide = render_section_slide(prs, page)
        else:
            slide = render_content_slide(prs, page, base_dir)
        if page["note"]:
            slide.notes_slide.notes_text_frame.text = page["note"]

    prs.save(out_path)
    print(f"{md_path} -> {out_path} ({len(pages)} slides)")


def main():
    lect_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "lect")
    md_files = sorted(glob.glob(os.path.join(lect_dir, "CH-*-Content.md")))
    for md_path in md_files:
        out_path = os.path.splitext(md_path)[0] + ".pptx"
        build_pptx(md_path, out_path)


if __name__ == "__main__":
    main()
